"""HTTP entry point for Snap to Intelligence.

Run locally with:
    uvicorn main:app --reload --port 8000
"""

from __future__ import annotations

from pathlib import Path

from dotenv import load_dotenv
from fastapi import BackgroundTasks, FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field as PydanticField
from sqlalchemy import func, select
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import adapters
import assemble
import mock_store
from database import Base, SessionLocal, engine
from workspace_models import ActivityLog, Field, Item, Workspace
from workspace_service import activity_data, field_data, item_data, run_item_pipeline, update_item_status


BASE_DIR = Path(__file__).resolve().parent
UPLOADS_DIR = BASE_DIR / "uploads"
UPLOADS_DIR.mkdir(exist_ok=True)
load_dotenv(BASE_DIR / ".env")
Base.metadata.create_all(bind=engine)

app = FastAPI(title="Snap to Intelligence API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000", "http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/api/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")

IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class WorkspaceCreate(BaseModel):
    name: str = PydanticField(min_length=1, max_length=200)
    created_by: str = PydanticField(default="local_user", max_length=100)


class ItemPatch(BaseModel):
    name: str | None = PydanticField(default=None, min_length=1, max_length=200)
    field_name: str | None = None
    action: str | None = None
    chosen_value: str | None = None
    actor: str = "local_user"


def response(body: dict, status_code: int = 200) -> JSONResponse:
    """Return the response format consumed by the UI on success and failure."""
    return JSONResponse(content=body, status_code=status_code)


def workspace_data(workspace: Workspace, item_count: int | None = None, include_items=False) -> dict:
    data = {"id": workspace.id, "name": workspace.name, "created_at": workspace.created_at.isoformat(),
            "created_by": workspace.created_by, "item_count": item_count if item_count is not None else len(workspace.items)}
    if include_items:
        data["items"] = [item_data(item) for item in sorted(workspace.items, key=lambda x: x.updated_at, reverse=True)]
    return data


@app.post("/workspaces", status_code=201)
def create_workspace(payload: WorkspaceCreate) -> dict:
    db = SessionLocal()
    try:
        workspace = Workspace(name=payload.name.strip(), created_by=payload.created_by)
        db.add(workspace); db.commit(); db.refresh(workspace)
        return workspace_data(workspace, item_count=0)
    finally:
        db.close()


@app.get("/workspaces")
def list_workspaces() -> list[dict]:
    db = SessionLocal()
    try:
        rows = db.execute(select(Workspace, func.count(Item.id)).outerjoin(Item).group_by(Workspace.id).order_by(Workspace.created_at.desc())).all()
        return [workspace_data(workspace, item_count=count) for workspace, count in rows]
    finally:
        db.close()


@app.get("/workspaces/{workspace_id}")
def get_workspace(workspace_id: int) -> JSONResponse:
    db = SessionLocal()
    try:
        workspace = db.get(Workspace, workspace_id)
        return response(workspace_data(workspace, include_items=True)) if workspace else response({"detail": "Workspace not found."}, 404)
    finally:
        db.close()


@app.post("/workspaces/{workspace_id}/items", status_code=202)
async def create_item(workspace_id: int, background_tasks: BackgroundTasks, image: UploadFile = File(...)) -> JSONResponse:
    if image.content_type not in IMAGE_TYPES:
        return response({"detail": "Upload a JPEG, PNG, or WebP image."}, 400)
    contents = await image.read()
    if not contents:
        return response({"detail": "The uploaded image is empty."}, 400)
    if len(contents) > MAX_UPLOAD_BYTES:
        return response({"detail": "Images must be 10 MB or smaller."}, 413)
    db = SessionLocal()
    try:
        if db.get(Workspace, workspace_id) is None:
            return response({"detail": "Workspace not found."}, 404)
        suffix = Path(image.filename or "image.jpg").suffix.lower() or ".jpg"
        stored_name = f"workspace-{assemble.new_profile_id()}{suffix}"
        stored_path = UPLOADS_DIR / stored_name
        stored_path.write_bytes(contents)
        item = Item(workspace_id=workspace_id, name=Path(image.filename or "Untitled item").stem,
                    image_urls=[f"/api/uploads/{stored_name}"], status="processing")
        db.add(item); db.flush()
        db.add(ActivityLog(item_id=item.id, workspace_id=workspace_id, actor="local_user", action="item_created",
                           new_value=item.name))
        db.commit(); db.refresh(item)
        background_tasks.add_task(run_item_pipeline, item.id, str(stored_path))
        return response(item_data(item), 202)
    finally:
        db.close()


@app.get("/workspaces/{workspace_id}/items")
def workspace_items(workspace_id: int) -> JSONResponse:
    db = SessionLocal()
    try:
        if db.get(Workspace, workspace_id) is None:
            return response({"detail": "Workspace not found."}, 404)
        items = db.scalars(select(Item).where(Item.workspace_id == workspace_id).order_by(Item.updated_at.desc())).all()
        return response({"items": [item_data(item) for item in items]})
    finally:
        db.close()


@app.get("/items/{item_id}")
def get_item(item_id: int) -> JSONResponse:
    db = SessionLocal()
    try:
        item = db.get(Item, item_id)
        return response(item_data(item, include_fields=True)) if item else response({"detail": "Item not found."}, 404)
    finally:
        db.close()


@app.patch("/items/{item_id}")
def patch_item(item_id: int, payload: ItemPatch) -> JSONResponse:
    db = SessionLocal()
    try:
        item = db.get(Item, item_id)
        if item is None:
            return response({"detail": "Item not found."}, 404)
        if payload.name is not None and payload.action is None:
            old_name, item.name = item.name, payload.name.strip()
            db.add(ActivityLog(item_id=item.id, workspace_id=item.workspace_id, actor=payload.actor, action="item_renamed", old_value=old_name, new_value=item.name))
        elif payload.field_name and payload.action in {"accept_current", "accept_alternate", "manual_edit"}:
            field = db.scalar(select(Field).where(Field.item_id == item.id, Field.field_name == payload.field_name))
            if field is None:
                return response({"detail": "Field not found on this item."}, 404)
            old_value, new_value = field.value, field.value
            if payload.action == "accept_alternate":
                alternates = field.conflicting_values or []
                matched = next((value for value in alternates if value == payload.chosen_value or (isinstance(value, dict) and str(value.get("value")) == payload.chosen_value)), None)
                if matched is None:
                    return response({"detail": "chosen_value must be one of the field's conflicting values."}, 400)
                new_value = matched.get("value") if isinstance(matched, dict) else matched
            elif payload.action == "manual_edit":
                if not payload.chosen_value or not payload.chosen_value.strip():
                    return response({"detail": "chosen_value is required for a manual edit."}, 400)
                new_value = payload.chosen_value.strip()
            field.value, field.confidence, field.is_user_verified, field.source_type, field.conflicting_values = new_value, 1.0, True, "user_verified", []
            db.add(ActivityLog(item_id=item.id, workspace_id=item.workspace_id, actor=payload.actor, action="field_updated", field_name=field.field_name, old_value=old_value, new_value=new_value))
            update_item_status(item)
        else:
            return response({"detail": "Send {name} or {field_name, action, chosen_value?, actor}."}, 400)
        db.commit(); db.refresh(item)
        return response(item_data(item, include_fields=True))
    finally:
        db.close()


@app.get("/workspaces/{workspace_id}/activity")
def workspace_activity(workspace_id: int) -> JSONResponse:
    db = SessionLocal()
    try:
        if db.get(Workspace, workspace_id) is None:
            return response({"detail": "Workspace not found."}, 404)
        logs = db.scalars(select(ActivityLog).where(ActivityLog.workspace_id == workspace_id).order_by(ActivityLog.timestamp.desc())).all()
        return response({"activity": [activity_data(log) for log in logs]})
    finally:
        db.close()


@app.get("/items/{item_id}/activity")
def item_activity(item_id: int) -> JSONResponse:
    db = SessionLocal()
    try:
        if db.get(Item, item_id) is None:
            return response({"detail": "Item not found."}, 404)
        logs = db.scalars(select(ActivityLog).where(ActivityLog.item_id == item_id).order_by(ActivityLog.timestamp.desc())).all()
        return response({"activity": [activity_data(log) for log in logs]})
    finally:
        db.close()


@app.get("/workspaces/{workspace_id}/export.xlsx")
def export_workspace(workspace_id: int):
    db = SessionLocal()
    try:
        workspace = db.get(Workspace, workspace_id)
        if workspace is None:
            return response({"detail": "Workspace not found."}, 404)
        items = db.scalars(select(Item).where(Item.workspace_id == workspace_id).order_by(Item.updated_at.desc())).all()
        workbook = Workbook(); summary = workbook.active; summary.title = "Summary"
        summary.append(["Name", "Brand", "Model Number", "Completeness Score", "Status", "Last Updated"])
        detail = workbook.create_sheet("Field Detail")
        detail.append(["Item Name", "Field Name", "Value", "Confidence", "Source Type", "Source Reference", "Verified by User"])
        header_fill = PatternFill("solid", fgColor="174B46")
        for sheet in (summary, detail):
            for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = header_fill
            sheet.freeze_panes = "A2"
        ready_fill, review_fill = PatternFill("solid", fgColor="E3F5E9"), PatternFill("solid", fgColor="FFF0CC")
        for item in items:
            fields = item.fields
            by_name = {field.field_name.lower(): field.value for field in fields}
            summary.append([item.name, by_name.get("brand", ""), by_name.get("model_number", by_name.get("model", "")), item.completeness_score, item.status, item.updated_at.isoformat()])
            fill = ready_fill if item.status == "ready" else review_fill if item.status == "needs_review" else None
            if fill:
                for cell in summary[summary.max_row]: cell.fill = fill
            for field in fields:
                detail.append([item.name, field.field_name, field.value, field.confidence, field.source_type, field.source_reference, "Yes" if field.is_user_verified else "No"])
                detail.cell(detail.max_row, 4).number_format = "0%"
        for sheet in (summary, detail):
            for column in sheet.columns:
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        buffer = BytesIO(); workbook.save(buffer); buffer.seek(0)
        filename = "".join(char if char.isalnum() else "_" for char in workspace.name).strip("_") or "workspace"
        return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}_export.xlsx"'})
    finally:
        db.close()


@app.get("/api/health")
def health() -> dict:
    return {
        "status": "ok",
        "extraction_module": adapters.extraction_available(),
        "enrichment_module": adapters.enrichment_available(),
        "mocks_loaded": mock_store.count(),
    }


@app.get("/api/mocks")
def list_mocks() -> dict:
    return {"available": mock_store.available(), "mocks": mock_store.summaries()}


@app.get("/api/mocks/{mock_id}")
def get_mock(mock_id: str) -> JSONResponse:
    item = mock_store.get(mock_id)
    if item is None:
        return response(assemble.failed("Sample profile was not found."), 404)
    return response(item)


@app.post("/api/scan")
async def scan(
    image: UploadFile | None = File(default=None),
    mode: str = Form(default="auto"),
) -> JSONResponse:
    """Extract a nameplate image, optionally returning a no-network sample."""
    if mode == "mock":
        sample = mock_store.first()
        if sample is None:
            return response(assemble.failed("No sample profiles are available."), 503)
        return response(sample)

    if mode not in {"auto", "extract_only"}:
        return response(assemble.failed("mode must be auto, extract_only, or mock."), 400)
    if image is None:
        return response(assemble.failed("Choose an image to scan."), 400)
    if image.content_type not in IMAGE_TYPES:
        return response(assemble.failed("Upload a JPEG, PNG, or WebP image."), 400)

    data = await image.read()
    if not data:
        return response(assemble.failed("The uploaded image is empty."), 400)
    if len(data) > MAX_UPLOAD_BYTES:
        return response(assemble.failed("Images must be 10 MB or smaller."), 413)

    suffix = Path(image.filename or "image.jpg").suffix.lower() or ".jpg"
    profile_id = assemble.new_profile_id()
    stored_name = f"{profile_id}{suffix}"
    stored_path = UPLOADS_DIR / stored_name
    stored_path.write_bytes(data)

    extraction = adapters.run_extraction(str(stored_path))
    if extraction.get("error") and not extraction.get("extracted_fields"):
        return response(assemble.failed(str(extraction["error"])), 502)

    enriched = None if mode == "extract_only" else adapters.run_enrichment(extraction)
    envelope = assemble.build_envelope(
        extraction=extraction,
        enriched=enriched,
        profile_id=profile_id,
        filename=image.filename or stored_name,
        image_url=f"/api/uploads/{stored_name}",
    )
    return response(envelope)
